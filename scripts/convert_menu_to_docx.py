#!/usr/bin/env python3
"""
Convert the JS menu data files into a DOCX (Microsoft Word) file.

Usage:
  python scripts/convert_menu_to_docx.py --out menu.docx

By default, it looks for the files in `src/data/hrana.js`, `src/data/drinks.js`,
and `src/data/pizzaData.js` relative to the repository root.
"""
import json
import re
from pathlib import Path
import argparse

from docx import Document
from docx.shared import Pt


def extract_js_array_text(text: str) -> str:
    start = text.find("[")
    if start == -1:
        raise ValueError("No array bracket '[' found in file content")
    end = text.rfind("]")
    if end == -1:
        raise ValueError("No closing array bracket ']' found in file content")
    return text[start:end+1]


def js_array_to_python(obj_text: str):
    s = obj_text
    s = re.sub(r"(\b\w+\b)\s*:", r'"\1":', s)
    s = s.replace("'", '"')
    s = re.sub(r",\s*([}\]])", r"\1", s)
    return json.loads(s)


def load_js_data(path: Path):
    text = path.read_text(encoding="utf-8")
    arr_text = extract_js_array_text(text)
    return js_array_to_python(arr_text)


def add_menu_item(doc, name, price, kol=None, jumbo=None):
    # Build the item line with name and price
    item_text = f"{name} ............ {price}"

    # Add extras on the same line if present
    extras = []
    if kol:
        extras.append(kol)
    if jumbo:
        extras.append(jumbo)
    if extras:
        item_text += f" ({', '.join(extras)})"

    # Add as a simple paragraph - fully editable
    doc.add_paragraph(item_text)


def generate_docx(hrana, drinks, pizza, out_path: Path, font_name: str = None):
    doc = Document()

    # set Normal style font if provided
    if font_name:
        style = doc.styles['Normal']
        style.font.name = font_name
        style.font.size = Pt(10)

    sections = [("Hrana", hrana), ("Drinks", drinks), ("Pizza", pizza)]

    for section_title, section_data in sections:
        doc.add_heading(section_title, level=1)
        for category in section_data:
            doc.add_heading(category.get('name', ''), level=2)
            for item in category.get('content', []):
                add_menu_item(doc, item.get('name', ''), item.get('price', ''), item.get('kol', ''), item.get('jumbo', ''))

            # small paragraph gap
            doc.add_paragraph()

    doc.save(str(out_path))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-dir", default="src/data", help="Path to data directory (default: src/data)")
    parser.add_argument("--out", default="menu.docx", help="Output DOCX file (default: menu.docx)")
    parser.add_argument("--font-name", default=None, help="Optional font name to use in DOCX (must be installed on system)")
    args = parser.parse_args()

    base = Path(__file__).resolve().parents[1]
    data_dir = (base / args.data_dir).resolve()

    hrana_path = data_dir / "hrana.js"
    drinks_path = data_dir / "drinks.js"
    pizza_path = data_dir / "pizzaData.js"

    if not hrana_path.exists() or not drinks_path.exists() or not pizza_path.exists():
        print("One or more data files missing in:", data_dir)
        print("Expected: hrana.js, drinks.js, pizzaData.js")
        return

    hrana = load_js_data(hrana_path)
    drinks = load_js_data(drinks_path)
    pizza = load_js_data(pizza_path)

    out_path = Path(args.out).resolve()
    generate_docx(hrana, drinks, pizza, out_path, font_name=args.font_name)
    print(f"DOCX written to {out_path}")


if __name__ == "__main__":
    main()
